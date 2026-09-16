# 임업진흥원 AI 상세페이지 지원사업 — 추진현황 보고서 만들기 (2026-09-16)
#   쓰기: python scripts/forest-report/build_report.py <데이터.json>
#        (생략하면 deliverables/forest-report/report-data.json 을 읽는다)
#   결과: deliverables/forest-report/<파일명>.docx  ·  같은 폴더에 _한글용.html(한글 변환용, EUC-KR)
#        한글(.hwp)·PDF 는 이어서 scripts/forest-report/to_hwp.ps1 로 만든다.
#
#   ▸ 문서 형식(6개 절)은 발주기관 제출용으로 사장님이 확정한 양식이다 — 절 순서·표 구성은 바꾸지 않는다.
#   ▸ 업체 응대 과정의 민감한 기록(전화 불통·불만·독촉 등)과 개인정보(성명·연락처·이메일·주소)는
#     데이터 단계에서 이미 빼거나 중립적으로 바꿔 둔다. 이 스크립트는 받은 내용을 그대로 옮긴다.
import json
import os
import sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", ".."))
OUT = os.path.join(ROOT, "deliverables", "forest-report")
os.makedirs(OUT, exist_ok=True)
DATA = sys.argv[1] if len(sys.argv) > 1 else os.path.join(OUT, "report-data.json")
D = json.load(open(DATA, encoding="utf-8"))

BODY = "맑은 고딕"
C = WD_ALIGN_PARAGRAPH.CENTER
doc = Document()

sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = sec.right_margin = Cm(2.0)
sec.top_margin = sec.bottom_margin = Cm(2.0)

st = doc.styles["Normal"]
st.font.name = BODY; st.font.size = Pt(10.5)
st.element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
st.paragraph_format.space_after = Pt(4)
st.paragraph_format.line_spacing = 1.45


def run(p, text, size=10.5, bold=False, color=None):
    r = p.add_run(text)
    r.font.name = BODY; r.font.size = Pt(size); r.bold = bold
    r._element.rPr.rFonts.set(qn("w:eastAsia"), BODY)
    if color:
        r.font.color.rgb = RGBColor(*color)
    return r


def para(text="", size=10.5, bold=False, align=None, before=0, after=4, color=None, indent=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if align is not None:
        p.alignment = align
    if indent is not None:
        p.paragraph_format.left_indent = Cm(indent)
        p.paragraph_format.first_line_indent = Cm(-0.45)
    if text:
        run(p, text, size, bold, color)
    return p


def heading(text):
    p = para(text, size=13, bold=True, before=16, after=6)
    pPr = p._element.get_or_add_pPr()
    bd = OxmlElement("w:pBdr"); bt = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "10"), ("w:space", "4"), ("w:color", "222222")):
        bt.set(qn(k), v)
    bd.append(bt); pPr.append(bd)


def bullet(label, text):
    p = para(indent=0.45, after=3)
    run(p, "· ", 10.5, True)
    if label:
        run(p, label + " ", 10.5, True)
    run(p, text, 10.5)


def shade(cell, color):
    sh = OxmlElement("w:shd")
    for k, v in (("w:val", "clear"), ("w:fill", color)):
        sh.set(qn(k), v)
    cell._tc.get_or_add_tcPr().append(sh)


def cell_text(cell, text, size=9, bold=False, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(1); p.paragraph_format.space_after = Pt(1)
    p.paragraph_format.line_spacing = 1.2
    if align is not None:
        p.alignment = align
    run(p, text, size, bold)


def table(cols, rows, header=True, center_idx=(), size=9, repeat_header=False):
    t = doc.add_table(rows=len(rows) + (1 if header else 0), cols=len(cols))
    t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    if header:
        for i, (name, w) in enumerate(cols):
            t.columns[i].width = Cm(w)
            cell_text(t.rows[0].cells[i], name, size, True, C); shade(t.rows[0].cells[i], "F2F3F6")
        if repeat_header:   # 쪽이 넘어가도 머리글 반복
            th = OxmlElement("w:tblHeader"); th.set(qn("w:val"), "true")
            t.rows[0]._tr.get_or_add_trPr().append(th)
    for r, row in enumerate(rows, start=1 if header else 0):
        t.rows[r]._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for i, val in enumerate(row):
            c = t.rows[r].cells[i]; c.width = Cm(cols[i][1])
            cell_text(c, val, size, False, C if i in center_idx else None)
    return t


M = D["meta"]
para(M["docno"], size=9, color=(0x6F, 0x77, 0x89), after=2)
for line in M["title"]:
    para(line, size=18, bold=True, align=C, after=0)
para(M["subtitle"], size=10.5, align=C, before=4, after=12, color=(0x3D, 0x44, 0x52))

t = doc.add_table(rows=len(M["info"]), cols=2); t.style = "Table Grid"; t.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate(M["info"]):
    t.rows[i].cells[0].width = Cm(3.2); t.rows[i].cells[1].width = Cm(13.8)
    cell_text(t.rows[i].cells[0], k, 10, True, C); shade(t.rows[i].cells[0], "F2F3F6")
    cell_text(t.rows[i].cells[1], v, 10)

heading("1. 사업 개요")
para(D["overview"]["text"], after=6)
for lab, txt in D["overview"]["bullets"]:
    bullet(lab, txt)

heading("2. 추진 실적 요약")
S = D["summary"]["counts"]
t2 = doc.add_table(rows=2, cols=len(S)); t2.style = "Table Grid"; t2.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, (k, v) in enumerate(S):
    t2.columns[i].width = Cm(17.0 / len(S))
    cell_text(t2.rows[0].cells[i], k, 9, True, C); shade(t2.rows[0].cells[i], "F2F3F6")
    cell_text(t2.rows[1].cells[i], v, 11, True, C)
para(after=4)
for lab, txt in D["summary"]["bullets"]:
    bullet(lab, txt)

heading("3. 업체별 추진 현황")
COLS = [(c[0], c[1]) for c in D["companies"]["cols"]]
table(COLS, D["companies"]["rows"], center_idx=tuple(D["companies"]["center"]), repeat_header=True)
para(D["companies"]["note"], size=9, after=2, color=(0x6F, 0x77, 0x89))

heading("4. 단계별 추진 경과")
for title_, body in D["steps"]:
    para(title_, size=10.5, bold=True, before=6, after=1)
    para(body, size=10.5, after=4, indent=0.45)

heading("5. 주요 사항 및 조치")
for lab, txt in D["issues"]:
    bullet(lab, txt)

heading("6. 향후 추진 계획")
table([(c[0], c[1]) for c in D["plan"]["cols"]], D["plan"]["rows"], center_idx=tuple(D["plan"]["center"]), size=9.5)

para(before=10)
t5 = doc.add_table(rows=1, cols=1); t5.style = "Table Grid"
shade(t5.rows[0].cells[0], "F7F8FA"); cell_text(t5.rows[0].cells[0], D["note"], 9.5)

para(before=18)
para(M["date"], size=11, align=C, after=6)
para(M["company"], size=14, bold=True, align=C)

docx_path = os.path.join(OUT, M["filename"] + ".docx")
doc.save(docx_path)

# ── 한글 변환용 HTML — 한글은 UTF-8 을 EUC-KR 로 읽으므로 cp949 로 저장한다 ──
def esc(s):
    return s.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")

h = ['<html><head><meta http-equiv="Content-Type" content="text/html; charset=euc-kr">',
     f'<title>{esc(M["subject"])}</title></head><body style="font-family:맑은 고딕; font-size:10.5pt">',
     f'<p align="left" style="font-size:9pt;color:#6f7789">{esc(M["docno"])}</p>',
     '<h1 style="text-align:center;font-size:17pt">' + "<br>".join(esc(x) for x in M["title"]) + "</h1>",
     f'<p style="text-align:center">{esc(M["subtitle"])}</p>',
     '<table border="1" cellspacing="0" cellpadding="5" width="100%">']
for k, v in M["info"]:
    h.append(f'<tr><td width="18%" bgcolor="#F2F3F6"><b>{esc(k)}</b></td><td>{esc(v)}</td></tr>')
h.append("</table>")


def sec_html(no_title, blocks):
    h.append(f'<h2 align="left" style="font-size:13pt">{esc(no_title)}</h2>')
    h.extend(blocks)


sec_html("1. 사업 개요", [f'<p align="left">{esc(D["overview"]["text"])}</p>'] +
         [f'<p align="left">· <b>{esc(l)}</b> {esc(t_)}</p>' for l, t_ in D["overview"]["bullets"]])

rows_html = ['<table border="1" cellspacing="0" cellpadding="5" width="100%"><tr>']
rows_html += [f'<td align="center" bgcolor="#F2F3F6"><b>{esc(k)}</b></td>' for k, _ in S]
rows_html.append("</tr><tr>")
rows_html += [f'<td align="center"><b>{esc(v)}</b></td>' for _, v in S]
rows_html.append("</tr></table>")
rows_html += [f'<p align="left">· <b>{esc(l)}</b> {esc(t_)}</p>' for l, t_ in D["summary"]["bullets"]]
sec_html("2. 추진 실적 요약", rows_html)

ch = ['<table border="1" cellspacing="0" cellpadding="4" width="100%" style="font-size:9pt"><tr>']
for c in D["companies"]["cols"]:
    ch.append(f'<td width="{c[2]}" align="center" bgcolor="#F2F3F6"><b>{esc(c[0])}</b></td>')
ch.append("</tr>")
for r in D["companies"]["rows"]:
    ch.append("<tr>" + "".join(
        f'<td valign="top" align="{"center" if i in D["companies"]["center"] else "left"}">{esc(v)}</td>'
        for i, v in enumerate(r)) + "</tr>")
ch.append("</table>")
ch.append(f'<p align="left" style="font-size:9pt;color:#6f7789">{esc(D["companies"]["note"])}</p>')
sec_html("3. 업체별 추진 현황", ch)

sec_html("4. 단계별 추진 경과", [f'<p align="left"><b>{esc(a)}</b><br>{esc(b)}</p>' for a, b in D["steps"]])
sec_html("5. 주요 사항 및 조치", [f'<p align="left">· <b>{esc(l)}</b> {esc(t_)}</p>' for l, t_ in D["issues"]])

ph = ['<table border="1" cellspacing="0" cellpadding="5" width="100%"><tr>']
ph += [f'<td align="center" bgcolor="#F2F3F6"><b>{esc(c[0])}</b></td>' for c in D["plan"]["cols"]]
ph.append("</tr>")
for r in D["plan"]["rows"]:
    ph.append("<tr>" + "".join(
        f'<td valign="top" align="{"center" if i in D["plan"]["center"] else "left"}">{esc(v)}</td>'
        for i, v in enumerate(r)) + "</tr>")
ph.append("</table>")
sec_html("6. 향후 추진 계획", ph)

h.append(f'<table border="1" cellspacing="0" cellpadding="8" width="100%"><tr><td bgcolor="#F7F8FA" style="font-size:9.5pt">{esc(D["note"])}</td></tr></table>')
h.append(f'<p style="text-align:center">{esc(M["date"])}</p>')
h.append(f'<p style="text-align:center;font-size:14pt"><b>{esc(M["company"])}</b></p></body></html>')

html_path = os.path.join(OUT, "_한글용.html")
# 한글은 줄바꿈이 CRLF 가 아니면 HTML 을 열지 못한다(2026-09-16 실측) — EUC-KR + CRLF 로 쓴다
open(html_path, "w", encoding="cp949", newline="\r\n").write("\n".join(h))

print(docx_path)
print(html_path)
